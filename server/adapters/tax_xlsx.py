"""Tax XLSX adapter — Choukalos Finances Summary workbook.

The workbook has many sheets; the one we pull is the **'Income'** sheet:
a header row with [Year, HHI, Taxes, Tax Rate%, Est Sav Amt, Est Gross Save%,
Est Spend, Est Spend%, <notes>] followed by one row per year (newest first).

Mapping (dollars -> integer cents):
  HHI         -> income_cents
  Taxes       -> income_tax_paid_cents
  Est Sav Amt -> savings_est_cents
  notes col   -> notes
savings_rate  = savings_est / income (recomputed for consistency)

The adapter locates the header row dynamically (row where col0=='Year' and
col1=='HHI'), so small layout shifts don't break it.
"""

import os

from .base import Adapter, LoadResult


class TaxXlsxAdapter(Adapter):
    name = "tax_xlsx"

    def detect(self, path: str) -> bool:
        return path.lower().endswith((".xlsx", ".xls"))

    def load(self, path: str) -> LoadResult:
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb["Income"] if "Income" in wb.sheetnames else None
        if ws is None:
            raise ValueError(f"{os.path.basename(path)}: no 'Income' sheet found")

        rows_iter = ws.iter_rows(values_only=True)
        header = None
        for row in rows_iter:
            if row and len(row) >= 5 and str(row[0]).strip() == "Year" and str(row[1]).strip() == "HHI":
                header = row
                break
        if header is None:
            raise ValueError(f"{os.path.basename(path)}: no Year/HHI header row found in 'Income' sheet")

        def col(name: str) -> int:
            for j, v in enumerate(header):
                if v is not None and str(v).strip() == name:
                    return j
            raise ValueError(f"column {name!r} not found in header {list(header)}")

        j_year, j_hhi, j_tax = col("Year"), col("HHI"), col("Taxes")
        j_sav = col("Est Sav Amt")
        j_notes = col("notes") if "notes" in [str(v).strip() for v in header if v is not None] else None

        result = LoadResult()
        for row in rows_iter:
            if row is None or j_year >= len(row):
                break
            raw_year = row[j_year]
            try:
                year = int(raw_year)
            except (TypeError, ValueError):
                break  # end of the yearly table
            if not (1900 <= year <= 2100):
                result.issues.append(f"skipped non-year value {raw_year!r}")
                continue

            def money(idx: int) -> int:
                v = row[idx] if idx < len(row) else None
                try:
                    return int(round(float(v) * 100))
                except (TypeError, ValueError):
                    result.issues.append(f"year {year}: bad money value {v!r} in col {idx}")
                    return 0

            income = money(j_hhi)
            tax = money(j_tax)
            sav = money(j_sav)
            notes = str(row[j_notes]).strip() if (j_notes is not None and j_notes < len(row) and row[j_notes]) else ""
            result.rows.append(
                {
                    "year": year,
                    "income_cents": income,
                    "income_tax_paid_cents": tax,
                    "savings_est_cents": sav,
                    "savings_rate": round(sav / income, 6) if income else None,
                    "notes": notes,
                }
            )
        if not result.rows:
            raise ValueError(f"{os.path.basename(path)}: no yearly rows found after header")
        return result